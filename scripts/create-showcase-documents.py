from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = ROOT / "public" / "demo-documents"
PURPLE = "541249"
LIGHT_PURPLE = "F8F0F6"
GRID = "D9D9D9"

PROJECTS = [
    {
        "slug": "sunu-energie",
        "company": "Sunu Énergie Décentralisée SA",
        "trade": "Sunu Énergie",
        "title": "Mini-réseaux solaires pour commerces de proximité",
        "country": "Côte d'Ivoire",
        "city": "Abidjan",
        "instrument": "Financement avec remboursement",
        "goal": "180 000 000 FCFA",
        "minimum": "250 000 FCFA",
        "terms": "11,5 % par an sur 30 mois, avec 3 mois de différé",
        "purpose": "Déployer 24 unités solaires avec stockage pour 310 commerces à Abidjan et Bouaké.",
        "budget": [
            ("Panneaux et onduleurs", "92 000 000 FCFA"),
            ("Batteries", "48 000 000 FCFA"),
            ("Installation", "24 000 000 FCFA"),
            ("Réserve opérationnelle", "16 000 000 FCFA"),
        ],
        "risks": "Variation du coût des équipements importés, rythme de raccordement et retards de paiement.",
    },
    {
        "slug": "naya-logistique",
        "company": "Naya Froid et Logistique SARL",
        "trade": "Naya Logistique",
        "title": "Plateforme frigorifique pour les filières locales",
        "country": "Sénégal",
        "city": "Dakar",
        "instrument": "Financement avec remboursement",
        "goal": "95 000 000 FCFA",
        "minimum": "100 000 FCFA",
        "terms": "10,25 % par an sur 24 mois, avec 2 mois de différé",
        "purpose": "Étendre une chaîne du froid entre producteurs, marchés urbains et restauration.",
        "budget": [
            ("Chambres froides", "41 000 000 FCFA"),
            ("Véhicules", "34 000 000 FCFA"),
            ("Suivi numérique", "8 000 000 FCFA"),
            ("Fonds de roulement", "12 000 000 FCFA"),
        ],
        "risks": "Saisonnalité des volumes, coût de l'énergie, maintenance et concentration initiale de la clientèle.",
    },
    {
        "slug": "kora-sante",
        "company": "Kora Santé Industries SAS",
        "trade": "Kora Santé",
        "title": "Unité régionale de consommables médicaux",
        "country": "Bénin",
        "city": "Cotonou",
        "instrument": "Ouverture du capital",
        "goal": "320 000 000 FCFA",
        "minimum": "500 000 FCFA",
        "terms": "18 % du capital proposé, valorisation avant opération de 1 450 000 000 FCFA",
        "purpose": "Installer une ligne de production, un laboratoire qualité et certifier les premiers produits.",
        "budget": [
            ("Ligne de production", "178 000 000 FCFA"),
            ("Laboratoire", "52 000 000 FCFA"),
            ("Certification", "28 000 000 FCFA"),
            ("Lancement et trésorerie", "62 000 000 FCFA"),
        ],
        "risks": "Délais de certification, montée en charge industrielle, prix des matières et absence de liquidité garantie.",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)


def set_run_font(run, name="Arial", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def build_docx(project):
    out_dir = OUTPUT_ROOT / project["slug"]
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "resume-du-projet.docx"

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(styles["Title"])
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11.5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Résumé du projet " + project["trade"])
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(project["title"])
    set_run_font(run, size=11.5, bold=True, color=PURPLE)

    notice = doc.add_paragraph()
    notice.paragraph_format.space_after = Pt(12)
    run = notice.add_run(
        "Document de démonstration. Toutes les données sont fictives et ne constituent ni une offre, ni une promesse de rendement."
    )
    set_run_font(run, size=9.5, color="555555")
    run.italic = True

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    run = intro.add_run(
        f"{project['company']} présente un besoin de financement de {project['goal']} à {project['city']}. "
        f"Le projet vise à {project['purpose'][0].lower() + project['purpose'][1:]}"
    )
    set_run_font(run)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_widths(table, [2.0, 4.2])
    headers = table.rows[0].cells
    headers[0].text = "Repère"
    headers[1].text = "Information"
    for cell in headers:
        set_cell_shading(cell, PURPLE)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    facts = [
        ("Entreprise", project["company"]),
        ("Implantation", f"{project['city']}, {project['country']}"),
        ("Mode de financement", project["instrument"]),
        ("Montant recherché", project["goal"]),
        ("Participation minimale", project["minimum"]),
        ("Conditions indicatives", project["terms"]),
    ]
    for index, (label, value) in enumerate(facts):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2:
                set_cell_shading(cell, "FAFAFA")
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.5, bold=(cell is cells[0]))

    doc.add_paragraph("Utilisation prévue des fonds", style="Heading 1")
    budget = doc.add_table(rows=1, cols=2)
    budget.autofit = False
    set_table_widths(budget, [4.1, 2.1])
    for i, text in enumerate(("Poste", "Montant indicatif")):
        budget.rows[0].cells[i].text = text
        set_cell_shading(budget.rows[0].cells[i], PURPLE)
        set_cell_borders(budget.rows[0].cells[i])
        for run in budget.rows[0].cells[i].paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    for index, (label, value) in enumerate(project["budget"]):
        cells = budget.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2:
                set_cell_shading(cell, LIGHT_PURPLE)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.5)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("Risques à examiner", style="Heading 1")
    risk = doc.add_paragraph(project["risks"])
    risk.paragraph_format.space_after = Pt(8)
    for run in risk.runs:
        set_run_font(run)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Dossier fictif de démonstration - septembre 2026")
    set_run_font(run, size=8.5, color="666666")

    doc.save(out_path)
    return out_path


def build_pdf(project):
    out_dir = OUTPUT_ROOT / project["slug"]
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "note-de-presentation.pdf"

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ProjectTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=19,
        leading=23,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=7,
    )
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#" + PURPLE),
        spaceAfter=10,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=13,
        textColor=colors.HexColor("#242124"),
        spaceAfter=7,
    )
    note_style = ParagraphStyle(
        "Note",
        parent=body_style,
        fontName="Helvetica-Oblique",
        fontSize=8.3,
        leading=11,
        textColor=colors.HexColor("#555555"),
    )
    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.black,
        spaceBefore=7,
        spaceAfter=5,
    )

    story = [
        Paragraph("Note de présentation " + project["trade"], title_style),
        Paragraph(project["title"], subtitle_style),
        Paragraph(
        "Document de démonstration. Toutes les données sont fictives et ne constituent ni une offre, ni une promesse de rendement.",
            note_style,
        ),
        Spacer(1, 5 * mm),
        Paragraph(
            f"{project['company']} présente un projet à {project['city']}. {project['purpose']} Le besoin indiqué est de {project['goal']}.",
            body_style,
        ),
    ]
    fact_data = [
        ["Repère", "Information"],
        ["Implantation", f"{project['city']}, {project['country']}"],
        ["Mode de financement", project["instrument"]],
        ["Participation minimale", project["minimum"]],
        ["Conditions indicatives", project["terms"]],
    ]
    facts = Table(fact_data, colWidths=[44 * mm, 123 * mm], repeatRows=1)
    facts.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + PURPLE)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("LEADING", (0, 0), (-1, -1), 11),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#" + GRID)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([facts, Paragraph("Utilisation prévue des fonds", heading_style)])
    budget_data = [["Poste", "Montant indicatif"], *[list(row) for row in project["budget"]]]
    budget = Table(budget_data, colWidths=[112 * mm, 55 * mm], repeatRows=1)
    budget.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + PURPLE)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#" + GRID)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + LIGHT_PURPLE)]),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend(
        [
            budget,
            Paragraph("Risques à examiner", heading_style),
            Paragraph(project["risks"], body_style),
        ]
    )

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=16 * mm,
        title="Note de présentation " + project["trade"],
        author="Équipe de démonstration",
    )
    doc.build(story)
    return out_path


if __name__ == "__main__":
    outputs = []
    for item in PROJECTS:
        outputs.append(build_docx(item))
        outputs.append(build_pdf(item))
    for path in outputs:
        print(path)
